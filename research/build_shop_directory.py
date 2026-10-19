"""Build the dated, manually researched workbook. Does not scrape or reverify data.
Run with openpyxl 3.1.5: python research/build_shop_directory.py
"""
from pathlib import Path
from datetime import datetime
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'deliverables' / 'Tariq_El_Jdideh_Shop_Directory_2026-10-16.xlsx'
CHECKED = datetime(2026, 10, 16)
# Excerpts are transcribed from retrieved pages or search results, not direct calls.
# An access date is not a business confirmation or a publisher's last-update date.
SOURCES = [
('S01', 'Najjar Stores', 'Business website', 'https://najjarstores.com/contact/', 'Page retrieved', 'Not stated on retrieved page', 'Kaskas St, Near TOTAL Station Beirut Lebanon; +961 71 496049 +961 01 666889; Monday – Saturday: 10:00 AM – 07:00 PM; Sunday: 10:00 AM – 01:00 PM', 'Name, phones, address, hours', 'Official contact page takes precedence over different directory hours.'),
('S02', 'Najjar Stores', 'Business social profile', 'https://www.facebook.com/NajjarStores/', 'Search-index text only', 'Not established', '01/666889 - 71/496049 (Whatsapp)', 'Phone corroboration and category', 'No direct business contact; social profile content seen through search.'),
('S03', 'Eldorado Shopping Mall', 'Business website', 'https://www.eldorado-mall.com/pages/stores.php', 'Page retrieved', 'Not stated on retrieved page', 'Tarik El Jadida; +9611858718; Mon-Sat 9:30 AM - 7:00 PM; Sunday 10:00 AM - 7:00 PM', 'Branch name, phone, area, hours', 'Branch-specific official store locator; no detailed street address given.'),
('S04', 'Makani Lebanon', 'Third-party directory', 'https://makanilebanon.com/portal/en-US/directory/group/eldorado-shopping-center/46/eldorado-shopping-center-tariq-el-jdideh/10421/', 'Search-index text only; direct fetch SSL error 525', 'Not established', 'Tareek El Jadideh > Al Malaab El Baladi; 01-858718; 01-858719; Monday to Saturday 9:00 to 20:00; Sunday 10:30 to 20:00', 'Eldorado landmark and older conflicting hours', 'Use S03 for phone/hours. Secondary number 01-858719 not included as confirmed current.'),
('S05', 'Yellow Pages Lebanon', 'Third-party directory', 'https://www.yellowpages.com.lb/en/business/clothing/eldorado-shopping-center', 'Search-index text only; direct fetch failed', 'Not established', 'Eldorado Shopping Center Tarik Jadideh, Tarik Jadida, Beirut, Lebanon; +9611858718', 'Eldorado branch and phone corroboration', 'Not the Airport Road or Hamra branch.'),
('S06', 'Halabi Bookshop', 'Business website', 'https://www.halabibookshop.com/contactus', 'Page retrieved', 'Not stated on retrieved page', 'Kaskas, El-Horge 57, Rue Jalloul 77, Nahhass Building, GF, Beirut, Lebanon; +961 1 851154; +961 71 501093', 'Name, phones, address', 'Use visible landline text: its underlying tel hyperlink incorrectly points to an unrelated US example number.'),
('S07', 'Halabi Bookshop', 'Business website', 'https://www.halabibookshop.com/', 'Page retrieved (both chunks)', 'Not stated; page includes 2026 content', 'Beirut, Kaskas, El-Horge 57, Jalloul Street 77, Nahhass Building, GF; WhatsApp at +96171501093', 'Address and WhatsApp corroboration', 'No opening schedule found in retrieved homepage/contact text.'),
('S08', 'Waze', 'Third-party map listing', 'https://www.waze.com/live-map/directions/lb/beirut-governorate/bayrut/halabi-bookshop?to=place.ChIJAXRzPmsXHxURFjju1g9ozgU', 'Search-index text only', 'Not established', '71 501 093; Monday 09:30 - 18:00; Tuesday 09:30 - 18:00; Wednesday 09:30 - 18:00; Thursday 09:30 - 18:00; Friday 09:30 - 18:00; Saturday 09:30 - 18:00', 'Halabi published schedule, not adopted as verified', 'Conflicts with S09; Sunday not explicitly shown in this result.'),
('S09', 'Wanderlog', 'Third-party map aggregator', 'https://wanderlog.com/place/details/3948634/halabi-bookshop', 'Search-index text only', 'Not established', 'Sunday: Closed; Monday: 9:30AM–3PM; Tuesday: 9:30AM–3PM; Wednesday: 9:30AM–3PM; Thursday: 9:30AM–3PM; Friday: 9:30AM–3PM; Saturday: 9:30AM–3PM; +961 71 501 093', 'Halabi conflicting published schedule', 'No current hours selected; ask shop directly.'),
('S10', 'Near Place', 'Third-party map directory', 'https://lb.near-place.com/shoppers-supermarket-tareek-el-jadida-tarik-al-jadida-beirut', 'Page retrieved (relevant first chunk)', 'Hours update not stated; displayed reviews 2017–2018', 'Tarik Al Jadida, Beirut; +961 1 700 717; Monday–Saturday 7:00 am - 9:00 pm; Sunday 7:00 am - 2:00 pm; Latitude: 33.8736433, Longitude: 35.5026908', 'Shopper name, phone, area, map pin, hours', 'Old review dates do not establish the hours update date.'),
('S11', 'VyMaps', 'Third-party map directory', 'https://vymaps.com/LB/Shopper-s-Supermarket--Tareek-El-Jadida-74888/', 'Page retrieved (relevant first chunk)', 'Added 2017-03-13; hours update not stated', 'Tarik Al Jadida, Beirut; +961 1 700 717; Monday–Saturday 7:00 AM – 9:00 PM; Sunday 7:00 AM – 2:00 PM', 'Shopper name, phone, area and hours corroboration', 'Same pin/reviews as Near Place: likely shared upstream data, not independent verification.'),
('S12', 'Kompass', 'Third-party business directory', 'https://www.kompass.com/z/lb/a/supermarket-and-hypermarket-operators/8140003/page-2/', 'Search-index text only', 'Not established', "Shopper's Supermarket Ground floor, Itani building, Tarik El-Jdideh sector, Beirut Tel.+961-1-700717", 'Shopper building-level address and phone', 'Directory evidence only; present occupancy not confirmed.'),
('S13', 'NoGarlicNoOnions', 'Editorial visit / review', 'https://www.nogarlicnoonions.com/safsouf-maamoul-madd-and-an-exceptional-knefeh/', 'Page retrieved; header street address from search-index text', 'Article dated 2015-04-23', 'Soleiman Boustani street, Tarik Jdideh, Facing Arab University, Beirut, Lebanon; +961 1 704 504', 'Safsouf location, landline and documented shop visit', 'Historical editorial source, not current operating-hours proof. Street detail appears in search result header.'),
('S14', 'Safsouf Sweets', 'Business social profile', 'https://www.instagram.com/safsouf.sweets/', 'Search-index text only; direct fetch HTTP 403', 'Not established', 'Delivering Arab Delights since 1930 Beirut, Lebanon For delivery orders WhatsApp 81 08 42 42', 'Safsouf delivery WhatsApp', 'Delivery number is not a substitute for a confirmed shop opening schedule.'),
('S15', 'Yellow Pages Lebanon', 'Third-party directory', 'https://www.yellowpages.com.lb/en/business/pastry-shops/safsouf-sweets', 'Search-index text only; direct fetch failed', 'Not established', '+9611704504; Mon-Sat 5:30 AM to 10:00 PM; Sun 5:30 AM to 8:00 PM', 'Safsouf phone and conflicting hours', 'Hours conflict with S16; not adopted as verified.'),
('S16', 'Wanderlog', 'Third-party map aggregator', 'https://wanderlog.com/place/details/4240907/safsouf-sweets', 'Search-index text only', 'Not established', 'Sunday–Saturday 6AM–9PM; +961 1 704 504', 'Safsouf phone and conflicting hours', 'Different from S15; no current schedule selected.'),
('S17', 'MiCare', 'Supplier reseller list', 'https://mi-care.net/resellers/', 'Page retrieved (relevant first chunk)', 'Not stated on retrieved page', 'Chicho Cell; Tarek Jdide; 961 3 502 259', 'Chicho name, area and phone', 'Supplier listing, not shop-owned source. No hours given.'),
('S18', 'Chicho Cell on OLX / dubizzle', 'Seller storefront', 'https://www.olx.com.lb/companies/ChichoCell-8544', 'Search-index text only', 'Relative ad ages shown; not reliable absolute update dates', 'Chicho Cell; Tariq el Jdideh, Beirut; Malaab, Beirut', 'Shop name and neighborhood corroboration', 'Ads corroborate seller/area, not phone, full address or hours.'),
('S19', 'Autoyas', 'Third-party business/social mirror', 'https://www.autoyas.com/LB/Beirut/810361852322348/Hariri-Car-audio-and-DVD-system', 'Page retrieved', 'Mirrored posts include 2026-09-06; hours update not stated', 'Cornish El Mazraa, Tarik Jdideh, Abou Chaker Street; 009613761946; Monday–Saturday 09:00 - 19:00; For info 03761946', 'Hariri name, phone, address and published hours', 'Sunday not listed, so not assumed closed. Several mirrored posts repeat the full phone; one post has a truncated typo.'),
('S20', 'Index of Lebanon', 'Third-party business directory', 'https://www.indexoflebanon.com/directory/pharmacies-in-lebanon/lebanon-company-in-tarik-eljdideh_al-labban-tarik-al-jadida.php', 'Page retrieved (relevant first chunk)', 'Actual business-detail update not stated', 'Kaskas Building, Main Road, Tarik Eljdideh, Beirut, Lebanon; Telephone: +9611315211', 'Al Labban name, phone, address', 'No usable business-specific hours in retrieved text. Generic directory schema hours are not accepted.'),
('S21', 'GLinkco', 'Third-party pharmacy directory', 'https://glinkco.com/en-us/name/al-labban-tarik-al-jadida-pharmacy-beyrouth-lebanon/', 'Search-index text only', 'Search metadata 2018-10-07; detail update unproven', 'Al Labban-Tarik Al Jadida Pharmacy; Beyrouth-Tarik Al Jadida; +9611315211', 'Al Labban phone and area corroboration', 'Not Al Labban Al Jadida in Hamra, which has another number.'),
('S22', 'WorldPlaces', 'Third-party directory', 'https://lebanon.worldplaces.me/review/93694239-rankoussi-tarik-jdideh.html', 'Search-index text only; direct fetch human-verification page', 'Not established', 'Rankoussi Tarik Jdideh; 01311772; طريق الجديدة شارع السبيل; Monday–Saturday 10:00 - 19:30; Sunday: Closed; 33.87191,35.50183', 'Rankoussi name, phone, street, schedule and map pin', 'Single directory source for detailed fields; kept in Unconfirmed leads.'),
('S23', 'Rankoussi Tarik Jdideh', 'Business social profile', 'https://www.facebook.com/people/Rankoussi-Tarik-Jdideh/100063696978325/', 'Search-index text only; direct fetch HTTP 403', 'Not established', 'شركة رنكوسي للأقمشة (برادي ومفروشات) والبياضات - فرع بيروت طريق جديدة', 'Rankoussi branch identity and category', 'Retrieved search text did not corroborate phone, street or hours.'),
('S24', 'GLinkco', 'Third-party pharmacy directory', 'https://glinkco.com/en-us/name/mayssa-saneh-pharmacy-tarik-al-jadida-beirut-lebanon/', 'Page retrieved (relevant first chunk)', 'Search metadata 2018-11-14; detail update unproven', 'Mayssa` Saneh Pharmacy; Tarik Al Jadida-Rue Loubos; +9611700906', 'Name, phone and street', 'Single publisher; no hours in retrieved listing. Kept in Unconfirmed leads.'),
('S25', 'GLinkco', 'Third-party pharmacy directory', 'https://glinkco.com/en-us/name/al-taware-pharmacy-tarik-al-jadida-beirut-lebanon/', 'Page retrieved (relevant first chunk)', 'Search metadata 2018-10-02; detail update unproven', 'Al Taware` Pharmacy; Tarik Al Jadida-Face De L`Hopital Makassed; +9611300981', 'Name, phone and landmark', 'Single publisher; no hours in retrieved listing. Kept in Unconfirmed leads.'),
('S26', 'Near Place', 'Third-party map directory', 'https://lb.near-place.com/dkan-abw-fwad-alkhrfan-tariq-el-jdideh-beirut', 'Page retrieved (relevant first chunk)', 'Hours update not stated; review shown from 2017', 'دكان ابو فواد الخرفان; Tariq El Jdideh, Beirut; +961 1 852 365; Monday and Sunday 8:00 am - 12:00 pm; Tuesday–Saturday closed; 33.8698738,35.4987815', 'Name, phone, area, pin and unusual published hours', 'Unusual two-day schedule lacks corroboration; withheld as usable hours. Kept in Unconfirmed leads.'),
('S27', 'ArabPlaces', 'Third-party map directory', 'https://lb.arabplaces.com/beirut/najjar-stores-45950', 'Search-index text only', 'Not established', '+961 71 496 049; Monday–Saturday 09:00-20:00; Sunday 10:00-13:00', 'Najjar conflicting third-party hours', 'Official S01 schedule is preferred and shown in the directory.'),
]

# Main entries have either an accessible shop-owned contact page, multiple
# publisher references, or dated business-post evidence. This is NOT a guarantee
# of live operating status, independent sources, working phones, or current hours.
SHOPS = [
 dict(id='B01', name='Najjar Stores', category='Electronics, gaming, toys and outdoor equipment', phone1='+961 1 666 889', phone2='+961 71 496 049', address='Kaskas Street, near TOTAL station, Kaskas / Tariq El-Jdideh, Beirut', hours='Mon–Sat 10:00–19:00; Sun 10:00–13:00', status='Business website: phone, address and hours', notes='Shop-owned contact page used. Other directories show Mon–Sat 09:00–20:00; official hours preferred. Phone 2 is identified as WhatsApp on business social profile. Not called.', ids=['S01','S02','S27'], phone_ids=['S01','S02'], address_ids=['S01'], hours_ids=['S01'], url='https://najjarstores.com/contact/'),
 dict(id='B02', name='Eldorado Shopping Mall — Tarik El Jadida', category='Clothing, shoes and household department store', phone1='+961 1 858 718', phone2='', address='Tarik El Jadida, Beirut; municipal stadium / Al Malaab El Baladi area (directory landmark)', hours='Mon–Sat 09:30–19:00; Sun 10:00–19:00', status='Business website: branch phone and hours', notes='Official branch locator S03 overrides conflicting directory schedules. Landmark is from S04, not the official locator. Exact street/building not established. Secondary directory phone 01-858719 omitted as not confirmed current.', ids=['S03','S04','S05'], phone_ids=['S03','S05'], address_ids=['S03','S04'], hours_ids=['S03'], url='https://www.eldorado-mall.com/pages/stores.php'),
 dict(id='B03', name='Halabi Bookshop', category='Bookshop', phone1='+961 1 851 154', phone2='+961 71 501 093', address='Kaskas, El-Horge 57, Rue Jalloul 77, Nahhass Building, ground floor, Beirut', hours='UNVERIFIED — conflicting published schedules; contact shop', status='Business website: phones and address; hours conflict', notes='Official contact text confirms both numbers and address. Landline tel hyperlink on site is erroneous; visible Lebanese number used. Waze: Mon–Sat 09:30–18:00; Wanderlog: Mon–Sat 09:30–15:00, Sun closed. Other listings also differ. No official hours found.', ids=['S06','S07','S08','S09'], phone_ids=['S06','S07'], address_ids=['S06','S07'], hours_ids=['S08','S09'], url='https://www.halabibookshop.com/contactus'),
 dict(id='B04', name="Shopper's Supermarket — Tareek El Jadida", category='Supermarket', phone1='+961 1 700 717', phone2='', address='Ground floor, Itani Building, Tarik El-Jdideh, Beirut (building from Kompass)', hours='Published, not directly confirmed: Mon–Sat 07:00–21:00; Sun 07:00–14:00', status='Directory-backed; likely shared hours data', notes='Near Place and VyMaps agree, but share pin/reviews and may copy the same upstream source. Kompass corroborates branch phone and adds building. Listing/review evidence dates back to 2017–2018; current operation and hours not established.', ids=['S10','S11','S12'], phone_ids=['S10','S11','S12'], address_ids=['S10','S12'], hours_ids=['S10','S11'], url='https://lb.near-place.com/shoppers-supermarket-tareek-el-jadida-tarik-al-jadida-beirut', coords='33.8736433,35.5026908'),
 dict(id='B05', name='Safsouf Sweets', category='Oriental sweets and pastry shop', phone1='+961 1 704 504', phone2='+961 81 084 242', address='Soleiman Boustani Street, Tariq El-Jdideh, near Beirut Arab University, Beirut (historical editorial address)', hours='UNVERIFIED — conflicting published schedules; contact shop', status='Landline corroborated; social-profile delivery phone; hours conflict', notes='Phone 2 is delivery WhatsApp from business Instagram search text only. Yellow Pages: Mon–Sat 05:30–22:00, Sun 05:30–20:00; Wanderlog: daily 06:00–21:00. No authoritative current hours found. Street from 2015 editorial search header; location not field-checked.', ids=['S13','S14','S15','S16'], phone_ids=['S13','S14','S15','S16'], address_ids=['S13'], hours_ids=['S15','S16'], url='https://www.instagram.com/safsouf.sweets/'),
 dict(id='B06', name='Chicho Cell', category='Mobile phones and accessories', phone1='+961 3 502 259', phone2='', address='Tariq El-Jdideh / Malaab, Beirut; exact street/building not established', hours='UNVERIFIED — no supported operating schedule found', status='Supplier-listed phone; seller storefront supports area', notes='Phone and area from MiCare reseller list. OLX seller storefront corroborates shop name and area, not phone/hours. No detailed street address adopted. Contact shop for current location and schedule.', ids=['S17','S18'], phone_ids=['S17'], address_ids=['S17','S18'], hours_ids=[], url='https://mi-care.net/resellers/'),
 dict(id='B07', name='Hariri Car Audio and DVD System', category='Car audio / electronics sales and installation', phone1='+961 3 761 946', phone2='', address='Corniche El Mazraa, Tarik Jdideh, Abou Chaker Street, Beirut', hours='Published, not directly confirmed: Mon–Sat 09:00–19:00; Sun not stated', status='Third-party mirror with dated business posts', notes='Autoyas mirrors posts through 2026-09-06 repeating 03 761946; directory lists schedule/address. One mirrored post has a truncated phone typo; full listing number and repeated full numbers agree. No independent confirmation; do not assume Sunday closed.', ids=['S19'], phone_ids=['S19'], address_ids=['S19'], hours_ids=['S19'], url='https://www.autoyas.com/LB/Beirut/810361852322348/Hariri-Car-audio-and-DVD-system', coords='33.8765,35.4998'),
 dict(id='B08', name='Al Labban — Tarik Al Jadida Pharmacy', category='Pharmacy', phone1='+961 1 315 211', phone2='', address='Kaskas Building, Main Road, Tarik El-Jdideh, Beirut', hours='UNVERIFIED — no usable business-specific schedule found', status='Phone and area corroborated by two directories', notes='Index of Lebanon gives building and number; GLinkco corroborates branch/number. Not the similarly named Hamra pharmacy. Directory structured-data schedules repeated across unrelated businesses were not treated as authentic hours. Not a confirmed duty/24-hour pharmacy.', ids=['S20','S21'], phone_ids=['S20','S21'], address_ids=['S20','S21'], hours_ids=[], url='https://www.indexoflebanon.com/directory/pharmacies-in-lebanon/lebanon-company-in-tarik-eljdideh_al-labban-tarik-al-jadida.php'),
]
LEADS = [
 dict(id='L01', name='Rankoussi — Tarik Jdideh', category='Fabrics, curtains and household linens', phone1='+961 1 311 772', phone2='', address='Al Sabil Street (شارع السبيل), Tariq El-Jdideh, Beirut', hours='UNVERIFIED listing: Mon–Sat 10:00–19:30; Sun closed', status='Detailed fields single-sourced; direct page blocked', notes='Business Facebook search result supports branch identity only. Phone, street and hours come solely from WorldPlaces search text; direct page blocked by verification challenge. Confirm all details before use.', ids=['S22','S23'], phone_ids=['S22'], address_ids=['S22'], hours_ids=['S22'], url='https://lebanon.worldplaces.me/review/93694239-rankoussi-tarik-jdideh.html', coords='33.87191,35.50183'),
 dict(id='L02', name='Mayssa Saneh Pharmacy', category='Pharmacy', phone1='+961 1 700 906', phone2='', address='Rue Loubos, Tarik Al Jadida, Beirut', hours='UNVERIFIED — no hours published in retrieved listing', status='Single older pharmacy-directory listing', notes='Only one publisher supports these fields; search metadata dates to 2018, not a verified current update. No shop-owned corroboration or confirmed current operation. Not a confirmed duty/24-hour pharmacy.', ids=['S24'], phone_ids=['S24'], address_ids=['S24'], hours_ids=[], url='https://glinkco.com/en-us/name/mayssa-saneh-pharmacy-tarik-al-jadida-beirut-lebanon/'),
 dict(id='L03', name='Al Taware Pharmacy', category='Pharmacy', phone1='+961 1 300 981', phone2='', address='Facing Makassed Hospital, Tarik Al Jadida, Beirut', hours='UNVERIFIED — no hours published in retrieved listing', status='Single older pharmacy-directory listing', notes='Only one publisher supports these fields; search metadata dates to 2018, not a verified current update. No shop-owned corroboration or confirmed current operation. Not a confirmed duty/24-hour pharmacy.', ids=['S25'], phone_ids=['S25'], address_ids=['S25'], hours_ids=[], url='https://glinkco.com/en-us/name/al-taware-pharmacy-tarik-al-jadida-beirut-lebanon/'),
 dict(id='L04', name='دكان ابو فواد الخرفان — Abou Fouad Al Khourfan shop', category='Grocery / minimarket', phone1='+961 1 852 365', phone2='', address='Tariq El-Jdideh, Beirut; exact street not published', hours='UNVERIFIED — unusual published schedule withheld; see notes', status='Single older map-directory listing', notes='Near Place reports Mon and Sun 08:00–12:00, Tue–Sat closed. This unusual schedule is not corroborated and must not be used as verified hours. One displayed review is from 2017. English name is a transliteration of the Arabic listing.', ids=['S26'], phone_ids=['S26'], address_ids=['S26'], hours_ids=['S26'], url='https://lb.near-place.com/dkan-abw-fwad-alkhrfan-tariq-el-jdideh-beirut', coords='33.8698738,35.4987815'),
]

NAVY = '16324F'
TEAL = '087F8C'
AMBER = 'FFF2CC'
PALE = 'E8F4F6'
WHITE = 'FFFFFF'


def table_sheet(wb, name, title, warning, headers, rows, widths, table_name):
    ws = wb.create_sheet(name)
    last = get_column_letter(len(headers))
    ws.merge_cells(f'A1:{last}1')
    ws['A1'] = title
    ws['A1'].font = Font(name='Calibri', size=18, bold=True, color=WHITE)
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f'A2:{last}2')
    ws['A2'] = warning
    ws['A2'].fill = PatternFill('solid', fgColor=AMBER)
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 32
    for j, value in enumerate(headers, 1):
        c = ws.cell(4, j, value)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=TEAL)
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[4].height = 32
    for i, row in enumerate(rows, 5):
        for j, value in enumerate(row, 1):
            c = ws.cell(i, j, value)
            c.font = Font(name='Calibri', size=11, color=NAVY)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(value, datetime):
                c.number_format = 'yyyy-mm-dd'
            elif isinstance(value, str):
                # Explicit text preserves +961 and prevents formula injection.
                c.data_type = 's'
                if value.startswith('https://'):
                    c.hyperlink = value
                    c.font = Font(name='Calibri', size=11, color='0563C1', underline='single')
        ws.row_dimensions[i].height = 110
    if rows:
        table = Table(displayName=table_name, ref=f'A4:{last}{len(rows)+4}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    for j, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = 'C5'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = 'Source review: 2026-10-16 | Not directly confirmed with businesses'
    ws.oddFooter.right.text = 'Page &P of &N'
    return ws


def directory_rows(shops):
    rows = []
    for shop in shops:
        coords = shop.get('coords')
        if coords:
            map_link = f'https://www.google.com/maps/search/?api=1&query={coords}'
        else:
            map_link = 'Not supplied — use source address; no verified pin'
        rows.append([shop['id'], shop['name'], shop['category'], shop['phone1'], shop['phone2'], shop['address'], shop['hours'], shop['status'], shop['notes'], ', '.join(shop['ids']), shop['url'], map_link, CHECKED, 'Not contacted'])
    return rows


def build():
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = 'Tariq El-Jdideh shop directory — source review 2026-10-16'
    wb.properties.subject = 'Public-source business directory with evidence and uncertainty labels'
    wb.properties.creator = 'Research compilation'
    wb.properties.description = 'Published information only; not a certified current directory. Includes unconfirmed leads separately.'
    readme_rows = [
        ('Purpose', 'Selected shops in Tariq El-Jdideh, Beirut, with published names, phone numbers, locations and operating hours where support was found. This is a researched sample, not an exhaustive census.'),
        ('IMPORTANT: verification limit', 'No calls, messages or site visits were made. No business is certified as currently open, reachable or operating the listed schedule. Source review confirms what a publication says, not the present real-world fact. The request for all details to be verified could not be fully satisfied from public web sources.'),
        ('Workbook contents', f'{len(SHOPS)} source-backed shops in Shop directory; {len(LEADS)} weaker records in Unconfirmed leads; {len(SOURCES)} source records; field-to-source mapping in Field evidence; uncompleted follow-up in Confirmation log.'),
        ('Source review date', '2026-10-16 (UTC). This is the date information was reviewed for this workbook, NOT the source update date or a direct business confirmation date. Re-running the script does not reverify the sources.'),
        ('Geographic scope', 'Tariq El-Jdideh / Tarik Al Jadida / Tareek El Jadideh variants; includes the adjacent Kaskas / El-Horge commercial streets identified in sources. Neighborhood boundaries are approximate and were not independently surveyed. Exact source address wording is retained where possible.'),
        ('Shop scope', 'Retail stores, groceries, sweets/pastries, mobile phones, car-audio retail and pharmacies. Banks, clinics, hospitals, rental businesses and general neighborhood social pages were excluded. Pharmacies here are not confirmed emergency/duty pharmacies.'),
        ('Main sheet threshold', 'Entries have an accessible business-owned contact page, multiple publisher references, or dated mirrored business-post evidence. This threshold is not a claim of independent corroboration of every field. Read each field source and row note.'),
        ('Unconfirmed leads threshold', 'A single older directory, blocked detailed source or unsupported phone/hours means a record is separated as a lead, not presented as fully verified. These may be useful for follow-up but should not be relied on without direct contact.'),
        ('Hours policy', '24-hour local Beirut time. Official branch-specific hours preferred over conflicting directories. If no authoritative schedule resolves a conflict, the usable hours cell says UNVERIFIED; competing schedules remain in notes/source ledger. No missing Sunday was assumed closed. Holidays, Ramadan and exceptional closures are not established.'),
        ('Published-only hours', 'Shopper and Hariri schedules are visibly labeled published, not directly confirmed. Rankoussi schedule is unverified and appears only as a lead. Shared upstream map data is not independent verification.'),
        ('Phone policy', 'Numbers use international +961 format and are stored as text. The domestic trunk zero is removed (01 becomes +961 1; 03 becomes +961 3). Numbers were not dial-tested. Mobile/WhatsApp purpose is stated only where published. An empty second-phone cell means none supported, not that no second number exists.'),
        ('Source access quality', 'Sources distinguish retrieved page text from search-index-only text, HTTP errors and blocked pages. Search-index evidence may be stale and is not treated as a successful live-page check. Source ages and known conflicts are recorded; dates of reviews or search metadata are not proof of detail freshness.'),
        ('Map policy', 'Only supplied source coordinates generate map links. They are published pins, not field-verified survey positions. Hariri pin comes from the Autoyas Directions link. No coordinates or precise streets were invented. Where no supported pin exists, use the address/source and contact the shop.'),
        ('How to audit', 'Filter Shop directory by evidence status. Use source IDs in each row, then Field evidence for phone/location/hours references, and Sources for clickable URLs, access limitations and evidence excerpts. The excerpt column is a short transcription/summary, not a full archived page.'),
        ('How to finish verification', 'Use Confirmation log to call or message each business. Confirm exact branch/address, every phone, all seven days, holiday exceptions and current operating status. Record contact date/method and the staff member or business account confirmation before marking confirmed. The prefilled log truthfully says Not contacted.'),
    ]
    ws = table_sheet(wb, 'Read me', 'TARIQ EL-JDIDEH | SHOP DIRECTORY', 'PUBLIC-SOURCE RESEARCH — NOT A FULLY VERIFIED CURRENT DIRECTORY. Read verification limits before use.', ['Topic','Explanation'], readme_rows, [32,125], 'ReadMeTable')
    for i in range(5, len(readme_rows)+5):
        ws.row_dimensions[i].height = 68 if i != 6 else 92
    ws.freeze_panes = 'B5'
    headers = ['ID','Shop name','Category','Phone 1 (+961)','Phone 2 (+961)','Location / published address','Operating hours (Beirut local time)','Evidence status — not live verification','Cautions / conflicts / follow-up','Source IDs','Main source URL','Published map pin (not field verified)','Sources reviewed (UTC)','Direct confirmation']
    widths = [9,34,27,22,22,49,49,42,85,24,49,47,18,21]
    table_sheet(wb, 'Shop directory', 'SOURCE-BACKED SHOP LISTINGS | 8 SHOPS', 'No shop contacted. UNVERIFIED means unresolved or missing; published-only hours may be outdated. Evidence is field-specific.', headers, directory_rows(SHOPS), widths, 'ShopsTable')
    table_sheet(wb, 'Unconfirmed leads', 'UNCONFIRMED LEADS | 4 SHOPS', 'Do not treat these as verified entries: phone, current location, operating status and hours require direct confirmation.', headers, directory_rows(LEADS), widths, 'LeadsTable')
    evidence = []
    for shop in SHOPS + LEADS:
        evidence.append([shop['id'],shop['name'],', '.join(shop['phone_ids']),', '.join(shop['address_ids']),', '.join(shop['hours_ids']) or 'None — not established',shop['status'],shop['notes']])
    table_sheet(wb, 'Field evidence', 'FIELD-BY-FIELD SOURCE MAPPING', 'Source IDs map to the Sources sheet. Multiple citations do not imply independent sources or present-day correctness.', ['Shop ID','Shop name','Phone source IDs','Location source IDs','Hours source IDs','Evidence status','Field cautions'], evidence, [10,36,27,27,29,46,105], 'EvidenceTable')
    source_rows = [[sid,publisher,kind,url,access,age,excerpt,fields,note,CHECKED] for sid,publisher,kind,url,access,age,excerpt,fields,note in SOURCES]
    table_sheet(wb, 'Sources', 'SOURCE LEDGER | AUDIT TRAIL', 'Short excerpts/transcriptions from pages or search-index text. Review date is not publisher update date. Errors and blocked pages are disclosed.', ['Source ID','Publisher / account','Source type','Source URL','Access / retrieval result','Source date / freshness limitation','Evidence excerpt / short transcription','Fields supported','Limitations / interpretation','Reviewed (UTC)'], source_rows, [11,29,31,66,52,52,105,43,85,18], 'SourcesTable')
    followups = [[s['id'],s['name'],s['phone1'],s['phone2'],'Not contacted','','','','','','',''] for s in SHOPS+LEADS]
    ws = table_sheet(wb, 'Confirmation log', 'DIRECT CONFIRMATION | FOLLOW-UP REGISTER', 'No outreach has been performed. Blank confirmation fields are intentional. Fill only after real business confirmation.', ['Shop ID','Shop name','Phone 1','Phone 2','Contact status','Date/time + timezone','Method / business account','Person / role (if provided)','Phones confirmed','Exact address confirmed','Full weekly hours + exceptions confirmed','Evidence / notes'], followups, [10,36,23,23,23,27,34,32,36,49,59,65], 'ConfirmationTable')
    validation = DataValidation(type='list', formula1='"Not contacted,Attempted - no reply,Partially confirmed,Confirmed,Closed / relocated"')
    ws.add_data_validation(validation)
    validation.add(f'E5:E{4+len(followups)}')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.active = 0
    wb.save(OUTPUT)
    validate(OUTPUT)
    print(OUTPUT)


def validate(path):
    all_shops = SHOPS + LEADS
    source_ids = {s[0] for s in SOURCES}
    assert len(source_ids) == len(SOURCES)
    assert len({s['id'] for s in all_shops}) == len(all_shops)
    for s in all_shops:
        for phone in (s['phone1'], s['phone2']):
            if phone:
                assert re.fullmatch(r'\+961 [13] \d{3} \d{3}|\+961 \d{2} \d{3} \d{3}', phone), phone
        for field in ('ids','phone_ids','address_ids','hours_ids'):
            assert set(s[field]) <= source_ids
    book = load_workbook(path)
    assert len(book.sheetnames) == 6
    assert book['Shop directory'].max_row == len(SHOPS)+4
    assert book['Unconfirmed leads'].max_row == len(LEADS)+4
    assert book['Sources'].max_row == len(SOURCES)+4
    for sheet in book:
        assert len(sheet.tables) == 1
        for row in sheet:
            for cell in row:
                assert cell.data_type != 'f', f'Unexpected formula: {sheet.title}!{cell.coordinate}'
    for name in ('Shop directory','Unconfirmed leads'):
        for row in book[name].iter_rows(min_row=5):
            assert row[3].data_type == 's'
            assert row[12].number_format == 'yyyy-mm-dd'
            assert row[13].value == 'Not contacted'
    book.close()

if __name__ == '__main__':
    build()
