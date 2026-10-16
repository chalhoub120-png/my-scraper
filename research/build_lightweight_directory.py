"""Build a compact real Excel workbook and a self-contained browser edition.

Install: python -m pip install -r research/requirements.txt
Run: python research/build_lightweight_directory.py
Input: existing Arabic workbook in deliverables/. No new verification is done.
Output: shops_ar_small.xlsx and index.html (no browser dependencies or network calls).
"""
from html import escape
from pathlib import Path
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'deliverables'
INPUT = FOLDER / 'Tariq_El_Jdideh_Shop_Directory_AR_2026-10-16.xlsx'
XLSX = FOLDER / 'shops_ar_small.xlsx'
HTML = FOLDER / 'index.html'
WARNING = ('مراجعة المصادر: 2026-10-16. لم يُجرَ اتصال أو زيارة للمتاجر. '
           'البيانات المنشورة ليست تأكيداً لاستمرار العمل أو صحة الأرقام والساعات حالياً. '
           'هذه نسخة مختصرة من 12 إدراجاً فقط، وليست قائمة المتاجر الـ667 التي لم تُرسل بعد.')
GROUPS = {
    'B01': 'إلكترونيات وهواتف وصوتيات سيارات',
    'B06': 'إلكترونيات وهواتف وصوتيات سيارات',
    'B07': 'إلكترونيات وهواتف وصوتيات سيارات',
    'B02': 'ملابس وأقمشة وأدوات منزلية',
    'L01': 'ملابس وأقمشة وأدوات منزلية',
    'B03': 'كتب ومكتبات',
    'B04': 'بقالة وسوبرماركت',
    'L04': 'بقالة وسوبرماركت',
    'B05': 'حلويات ومعجنات',
    'B08': 'صيدليات', 'L02': 'صيدليات', 'L03': 'صيدليات',
}


def text(value):
    return escape(str(value or ''), quote=True)


def safe_url(value):
    # No executable schemes are allowed, even if the input workbook changes.
    url = str(value or '')
    if not url.startswith('https://'):
        raise ValueError(f'Unsupported source URL: {url!r}')
    return text(url)


def read_rows(sheet):
    return list(sheet.iter_rows(min_row=5, values_only=True))


def build_excel(shops, sources, shop_headers, source_headers):
    book = Workbook()
    book.remove(book.active)
    for name, headers, rows, widths in (
        ('المتاجر', [*shop_headers, 'تصنيف السجل'], shops,
         [10,35,28,23,23,48,48,42,84,24,48,48,20,22,29]),
        ('المصادر', source_headers, sources,
         [11,30,32,65,49,49,100,45,82,20]),
    ):
        sheet = book.create_sheet(name)
        sheet.sheet_view.rightToLeft = True
        sheet.sheet_view.showGridLines = False
        last = get_column_letter(len(headers))
        sheet.merge_cells(f'A1:{last}1')
        sheet.cell(1, 1, 'دليل متاجر طريق الجديدة — نسخة Excel مختصرة')
        sheet.merge_cells(f'A2:{last}2')
        sheet.cell(2, 1, WARNING + ' حُذفت أوراق الأدلة المكررة وسجل المتابعة الفارغ؛ حُفظت بيانات المتاجر وسجل المصادر كاملاً.')
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[2].height = 48
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(4, col, header)
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='16324F')
            cell.alignment = Alignment(wrap_text=True, horizontal='right')
        sheet.row_dimensions[4].height = 45
        for row_number, values in enumerate(rows, 5):
            for col, value in enumerate(values, 1):
                if value is None:
                    continue
                cell = sheet.cell(row_number, col, value)
                if isinstance(value, str):
                    cell.data_type = 's'  # Preserve phone numbers; never create formulas.
                if name == 'المتاجر' and col == 13 or name == 'المصادر' and col == 10:
                    cell.number_format = 'yyyy-mm-dd'
                ltr = str(value).startswith(('https://', '+961', 'S0', 'S1', 'S2', 'B0', 'L0'))
                cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left' if ltr else 'right', readingOrder=1 if ltr else 2)
                if str(value).startswith('https://'):
                    cell.hyperlink = str(value)
                    cell.font = Font(color='0563C1', underline='single')
            sheet.row_dimensions[row_number].height = 130
        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = 'C5'
        sheet.auto_filter.ref = f'A4:{last}{sheet.max_row}'
    book.save(XLSX)
    book.close()


def build_html(shops, sources):
    categories = list(dict.fromkeys(GROUPS.values()))
    options = ''.join(f'<option>{text(c)}</option>' for c in categories)
    sections = []
    for category in categories:
        cards = []
        for row in shops:
            sid, name, kind, phone1, phone2, address, hours, status, notes, refs, url, pin, reviewed, direct, tier = row
            if GROUPS[sid] != category:
                continue
            phones = []
            for phone in (phone1, phone2):
                if phone:
                    tel = str(phone).replace(' ', '')
                    phones.append(f'<a href="tel:{text(tel)}"><bdi dir="ltr">{text(phone)}</bdi></a>')
            citations = ' '.join(f'<a href="#{text(ref.strip())}">{text(ref.strip())}</a>' for ref in refs.split(','))
            map_link = f' · <a href="{safe_url(pin)}" target="_blank" rel="noopener noreferrer">نقطة الخريطة المنشورة</a>' if str(pin).startswith('https://') else ''
            badge = 'lead' if str(sid).startswith('L') else 'backed'
            cards.append(
                f'<article data-category="{text(category)}"><h3>{text(name)}</h3>'
                f'<span class="badge {badge}">{text(tier)}</span>'
                f'<p class="kind">{text(kind)}</p>'
                f'<dl><dt>الهاتف</dt><dd>{" · ".join(phones)}</dd>'
                f'<dt>الموقع</dt><dd>{text(address)}</dd>'
                f'<dt>ساعات العمل</dt><dd>{text(hours)}</dd>'
                f'<dt>حالة المعلومات</dt><dd>{text(status)}</dd></dl>'
                f'<details><summary>التنبيهات وحدود التحقق</summary><p>{text(notes)}</p>'
                '<p>التأكيد المباشر: لم يُجرَ تواصل.</p></details>'
                f'<p class="links"><a href="{safe_url(url)}" target="_blank" rel="noopener noreferrer">المصدر الرئيسي</a>{map_link} · المراجع: {citations}</p>'
                '</article>'
            )
        sections.append(f'<section class="category"><h2>{text(category)}</h2><div class="cards">{"".join(cards)}</div></section>')
    source_items = []
    for sid, publisher, kind, url, access, age, excerpt, fields, note, reviewed in sources:
        # Original-language excerpts stay in the Excel file; only the translated
        # provenance and limitations are duplicated in HTML to reduce its size.
        source_items.append(
            f'<li id="{text(sid)}"><b>{text(sid)} · {text(publisher)}</b> — {text(kind)} '
            f'<a href="{safe_url(url)}" target="_blank" rel="noopener noreferrer">فتح المصدر</a>'
            f'<p>{text(access)}؛ {text(age)}.</p><p>{text(fields)}. {text(note)}</p></li>'
        )
    document = '''<!doctype html><html lang="ar" dir="rtl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>دليل متاجر طريق الجديدة</title>
<style>
*{box-sizing:border-box}body{margin:0;background:#f3f6fa;color:#193047;font:16px/1.7 Arial,sans-serif}main{max-width:1150px;margin:auto;padding:24px}header{border-bottom:3px solid #087f8c;padding-bottom:16px}h1{font-size:28px;margin:0}h2{font-size:22px;margin:28px 0 12px}h3{font-size:19px;margin:0 0 10px}p{margin:8px 0}a{color:#065c9c;text-underline-offset:3px}.warning{background:#fff3ce;border-right:4px solid #b16a00;padding:14px;margin:18px 0}.toolbar{display:flex;gap:12px;flex-wrap:wrap;padding:14px 0}label{display:flex;flex-direction:column;gap:5px;flex:1;min-width:210px}input,select{font:inherit;padding:10px;border:1px solid #a9bac9;border-radius:6px;background:white;min-width:0;width:100%}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,430px),1fr));gap:16px}article{background:white;padding:20px;border:1px solid #d9e1e8;border-radius:9px}.kind,.links{font-size:14px}.badge{font-size:13px;padding:4px 8px;border-radius:5px;display:inline-block}.backed{background:#e3f2f2}.lead{background:#fff0d5}dl{margin:12px 0}dt{font-weight:bold;margin-top:8px}dd{margin:0}details{border-top:1px solid #e4e9ef;padding-top:10px}summary{cursor:pointer;font-weight:bold}.links{overflow-wrap:anywhere}.sources{background:white;padding:20px;border-radius:9px;margin-top:28px}.sources li{padding:12px 0;border-bottom:1px solid #ddd;scroll-margin-top:20px}.sources p{font-size:14px}footer{font-size:14px;margin-top:24px}[hidden]{display:none!important}:focus-visible{outline:3px solid #087f8c;outline-offset:3px}@media print{body{background:white}.toolbar{display:none}main{max-width:none;padding:0}.cards{display:block}article{break-inside:avoid;margin-bottom:12px}details::details-content{content-visibility:visible}}
</style><main><header><h1>دليل متاجر طريق الجديدة</h1><p>نسخة عربية خفيفة تعمل مباشرة في المتصفح، دون تثبيت برامج أو تحميل مكتبات.</p></header>
<div class="warning">__WARNING__</div><p>8 إدراجات مدعومة بمصادر و4 إدراجات غير مؤكدة. التصنيف يعني تنظيم المعلومات، وليس تأكيد صحتها الحالية. الساعات بتوقيت بيروت وبصيغة 24 ساعة.</p>
<div class="toolbar" id="filters" hidden><label>البحث بالاسم أو الهاتف أو الموقع<input id="q" type="search" placeholder="ابحث ضمن المتاجر…"></label><label>نوع النشاط<select id="category"><option value="">كل الأنشطة</option>__OPTIONS__</select></label></div><p id="count" role="status" aria-live="polite">عرض 12 من 12 إدراجاً</p><noscript><p>تظهر القائمة كاملة دون جافاسكربت. يلزم تفعيله لاستخدام البحث والتصفية فقط.</p></noscript>
__SECTIONS__
<section class="sources"><h2>المصادر وحدود التحقق</h2><p>هذه المراجع من البحث الأصلي، ولم يُجرَ تحقق جديد عند إنشاء هذه النسخة. المقتطفات بلغتها الأصلية محفوظة في ملف Excel؛ اختُصرت هنا لتقليل الحجم. يتضمن النطاق شوارع تجارية مجاورة في قصقص والحرج؛ لم تُمسح الحدود ميدانياً.</p><ol>__SOURCES__</ol></section><footer>الأسماء العربية المنقولة عن اللاتينية ليست بالضرورة الأسماء الرسمية. الصيدليات ليست مؤكدة كمراكز مناوبة أو خدمة على مدار الساعة. يُنصح بالاتصال قبل الزيارة.</footer></main>
<script>
'use strict';
const cards=Array.from(document.querySelectorAll('article'));
const groups=Array.from(document.querySelectorAll('.category'));
const query=document.getElementById('q'),category=document.getElementById('category');
function normalize(s){return s.toLocaleLowerCase().normalize('NFKD').replace(/[\\u064b-\\u065f\\u0670\\u0640]/g,'').replace(/[أإآ]/g,'ا').replace(/ى/g,'ي').replace(/\\s+/g,'');}
const searchText=cards.map(c=>normalize(c.textContent));
function filter(){const q=normalize(query.value);let shown=0;cards.forEach((c,i)=>{c.hidden=!((!category.value||c.dataset.category===category.value)&&(!q||searchText[i].includes(q)));if(!c.hidden)shown++;});groups.forEach(g=>g.hidden=!Array.from(g.querySelectorAll('article')).some(c=>!c.hidden));document.getElementById('count').textContent=`عرض ${shown} من ${cards.length} إدراجاً`;}
query.addEventListener('input',filter);category.addEventListener('change',filter);document.getElementById('filters').hidden=false;
</script></html>'''
    for key, value in {
        '__WARNING__': text(WARNING), '__OPTIONS__': options,
        '__SECTIONS__': ''.join(sections), '__SOURCES__': ''.join(source_items),
    }.items():
        document = document.replace(key, value)
    HTML.write_text(document, encoding='utf-8')


def main():
    source_book = load_workbook(INPUT)
    main_rows = read_rows(source_book['دليل المتاجر'])
    lead_rows = read_rows(source_book['إدراجات غير مؤكدة'])
    shops = [(*r, 'مدعوم بمصادر — غير مؤكد مباشرة') for r in main_rows]
    shops += [(*r, 'إدراج غير مؤكد') for r in lead_rows]
    sources = read_rows(source_book['المصادر'])
    shop_headers = [c.value for c in source_book['دليل المتاجر'][4]]
    source_headers = [c.value for c in source_book['المصادر'][4]]
    assert len(shops) == 12 and len(sources) == 27
    assert set(GROUPS) == {r[0] for r in shops}
    build_excel(shops, sources, shop_headers, source_headers)
    build_html(shops, sources)
    validate(shops, sources)
    source_book.close()
    for path in (XLSX, HTML):
        print(f'{path.relative_to(ROOT)}: {path.stat().st_size:,} bytes')
    reduction = 100 * (1 - XLSX.stat().st_size / INPUT.stat().st_size)
    print(f'Excel size reduction from full Arabic edition: {reduction:.1f}%')


def validate(shops, sources):
    from html.parser import HTMLParser
    class AuditHTML(HTMLParser):
        def __init__(self):
            super().__init__()
            self.cards = 0
            self.ids = set()
            self.refs = set()
        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            if tag == 'article':
                self.cards += 1
            if 'id' in attrs:
                assert attrs['id'] not in self.ids
                self.ids.add(attrs['id'])
            if tag == 'a' and attrs.get('href', '').startswith('#'):
                self.refs.add(attrs['href'][1:])
            assert not (tag == 'script' and 'src' in attrs)
    with ZipFile(XLSX) as archive:
        assert archive.testzip() is None
    result = load_workbook(XLSX)
    assert result.sheetnames == ['المتاجر', 'المصادر']
    assert read_rows(result['المتاجر']) == shops
    assert read_rows(result['المصادر']) == sources
    for sheet in result:
        assert sheet.sheet_view.rightToLeft
        assert all(c.data_type != 'f' for row in sheet for c in row)
    result.close()
    content = HTML.read_text(encoding='utf-8')
    parser = AuditHTML()
    parser.feed(content)
    assert parser.cards == len(shops)
    assert parser.refs <= parser.ids
    assert {r[0] for r in sources} <= parser.ids
    for row in shops:
        for value in (row[1], row[3], row[5], row[6], row[7], row[8]):
            assert text(value) in content
    assert '__WARNING__' not in content
    assert XLSX.stat().st_size < INPUT.stat().st_size
    print('PASS: shop/source data preserved; HTML references and XLSX integrity checked.')


if __name__ == '__main__':
    main()
