"""Create a plain .xlsx edition without tables, validation or custom metadata."""
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
source = root / 'deliverables/Tariq_El_Jdideh_Shop_Directory_AR_2026-10-16.xlsx'
output = root / 'deliverables/shops_ar.xlsx'
src = load_workbook(source)
out = Workbook()
out.remove(out.active)
for source_sheet in src:
    sheet = out.create_sheet(source_sheet.title)
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = 'C5' if source_sheet.max_column > 2 else 'B5'
    for row in source_sheet:
        for old in row:
            if old.value is None:
                continue
            cell = sheet.cell(old.row, old.column, old.value)
            if isinstance(old.value, str):
                cell.data_type = 's'
            cell.number_format = old.number_format
            cell.font = Font(name='Arial', size=11)
            ltr = str(old.value).startswith(('https://', '+961', 'S0', 'S1', 'S2', 'B0', 'L0'))
            cell.alignment = Alignment(horizontal='left' if ltr else 'right', vertical='top', wrap_text=True, readingOrder=1 if ltr else 2)
            if old.hyperlink:
                cell.hyperlink = old.hyperlink.target
                cell.font = Font(name='Arial', size=11, color='0563C1', underline='single')
    last = get_column_letter(source_sheet.max_column)
    for row in (1,2):
        sheet.merge_cells(f'A{row}:{last}{row}')
    sheet.row_dimensions[1].height = 36
    sheet.row_dimensions[2].height = 45
    sheet.row_dimensions[4].height = 48
    for cell in sheet[4]:
        cell.fill = PatternFill('solid', fgColor='16324F')
        cell.font = Font(name='Arial', size=11, color='FFFFFF', bold=True)
    for row in range(5,sheet.max_row+1):
        sheet.row_dimensions[row].height = source_sheet.row_dimensions[row].height or 110
    for col in range(1,source_sheet.max_column+1):
        letter = get_column_letter(col)
        sheet.column_dimensions[letter].width = source_sheet.column_dimensions[letter].width
    sheet.auto_filter.ref = f'A4:{last}{sheet.max_row}'
    sheet.sheet_view.showGridLines = False
out.active = 1
out.save(output)
with ZipFile(output) as archive:
    assert archive.testzip() is None
    assert '[Content_Types].xml' in archive.namelist()
    for entry in archive.namelist():
        if entry.endswith('.xml'):
            ET.fromstring(archive.read(entry))
check = load_workbook(output)
assert check.sheetnames == src.sheetnames
for before, after in zip(src,check):
    for row in before:
        for cell in row:
            assert after.cell(cell.row,cell.column).value == cell.value
    assert not after.tables
    assert not after.data_validations.dataValidation
check.close()
src.close()
print(f'Validated Excel .xlsx: {output} ({output.stat().st_size:,} bytes)')
